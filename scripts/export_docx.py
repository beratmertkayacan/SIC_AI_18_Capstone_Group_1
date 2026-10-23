"""Markdown raporlarını Word (.docx) dosyasına çevirir.

Pandoc ile dönüştürür, sonra python-docx ile yazı tipi, başlık rengi,
tablo kenarlıkları ve görsel genişliği gibi biçim ayarlarını yapar.

Kullanım (repo kök klasöründe):
    python scripts/export_docx.py docs/04_concept_note.md

Çıktı aynı klasöre aynı adla .docx olarak yazılır. Pandoc kurulu olmalıdır
(macOS: brew install pandoc).
"""

import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x1F, 0x5F, 0x6B)
HEADER_FILL = "E3EEF0"
BORDER_COLOR = "B7C4C7"
MAX_IMAGE_WIDTH = Cm(16.6)


def set_page(doc):
    """A4 sayfa ve kenar boşlukları."""
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2)


def set_styles(doc):
    """Gövde metni ve başlık stilleri."""
    styles = {style.name.lower(): style for style in doc.styles}

    for name in ["normal", "body text", "first paragraph", "compact"]:
        if name in styles:
            style = styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(10.5)
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Calibri")
            style.paragraph_format.space_after = Pt(6)

    for level, size in [(1, 18), (2, 14), (3, 12)]:
        name = f"heading {level}"
        if name in styles:
            style = styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = TEAL
            style.paragraph_format.space_before = Pt(18 if level == 1 else 14)
            style.paragraph_format.space_after = Pt(6)


def style_table(table):
    """İnce kenarlık, tam sayfa genişliği, renkli başlık satırı."""
    props = table._tbl.tblPr

    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), BORDER_COLOR)
        borders.append(line)
    for old in props.findall(qn("w:tblBorders")):
        props.remove(old)
    props.append(borders)

    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    for old in props.findall(qn("w:tblW")):
        props.remove(old)
    props.append(width)

    for row_index, row in enumerate(table.rows):
        is_header = row_index == 0
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    if is_header:
                        run.font.bold = True
            if is_header:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:val"), "clear")
                shade.set(qn("w:color"), "auto")
                shade.set(qn("w:fill"), HEADER_FILL)
                cell._tc.get_or_add_tcPr().append(shade)


def fit_images(doc):
    """Sayfadan taşan görselleri orantılı küçültür."""
    for shape in doc.inline_shapes:
        if shape.width > MAX_IMAGE_WIDTH:
            ratio = MAX_IMAGE_WIDTH / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)


def main():
    source = Path(sys.argv[1])
    output = source.with_suffix(".docx")
    subprocess.run(
        ["pandoc", str(source), "-f", "gfm", "-o", str(output), f"--resource-path={source.parent}"],
        check=True,
    )

    doc = Document(output)
    set_page(doc)
    set_styles(doc)
    for table in doc.tables:
        style_table(table)
    fit_images(doc)
    doc.save(output)
    print(f"Yazıldı: {output}")


if __name__ == "__main__":
    main()
